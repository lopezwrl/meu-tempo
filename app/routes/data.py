"""Dados: backup completo (JSON), restauração e exportações (CSV / Excel)."""
import csv
import io
import os
import re
import shutil
from datetime import date, datetime, time as dtime
from urllib.parse import quote

from flask import Blueprint, Response, jsonify, request
from sqlalchemy import Date, DateTime, Time

from app.models.db import db
from app.models.category import Category
from app.models.project import Project
from app.models.task import Task
from app.models.time_entry import TimeEntry
from app.models.user import User
from app.utils.current_user import CURRENT_USER_ID

data_bp = Blueprint("data", __name__)
BACKUP_FORMAT = 1


def _attachment(data, mimetype, filename):
    resp = Response(data, mimetype=mimetype)
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "-", filename)
    resp.headers["Content-Disposition"] = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
    return resp


def _stamp():
    return datetime.now().strftime("%Y%m%d-%H%M")


# ---------------- Backup / restauração ----------------
def _serialize(value):
    if isinstance(value, (datetime, date, dtime)):
        return value.isoformat()
    return value


def _parse(column, value):
    if value is None:
        return None
    if isinstance(column.type, DateTime):
        return datetime.fromisoformat(value)
    if isinstance(column.type, Date):
        return date.fromisoformat(value)
    if isinstance(column.type, Time):
        return dtime.fromisoformat(value)
    return value


@data_bp.route("/api/backup")
def api_backup():
    payload = {
        "app": "meu-tempo",
        "format": BACKUP_FORMAT,
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "tables": {},
    }
    for table in db.metadata.sorted_tables:
        rows = db.session.execute(table.select()).mappings().all()
        payload["tables"][table.name] = [{k: _serialize(v) for k, v in row.items()} for row in rows]
    import json
    return _attachment(json.dumps(payload, ensure_ascii=False, indent=1), "application/json", f"meu-tempo-backup-{_stamp()}.json")


@data_bp.route("/api/backup/restore", methods=["POST"])
def api_restore():
    """Substitui TODOS os dados pelos do arquivo de backup (guarda uma cópia do banco antes)."""
    payload = request.get_json(silent=True) or {}
    tables = payload.get("tables")
    if payload.get("app") != "meu-tempo" or not isinstance(tables, dict):
        return jsonify({"error": "Este arquivo não parece um backup do Meu Tempo."}), 400
    if payload.get("format", 0) > BACKUP_FORMAT:
        return jsonify({"error": "Backup de uma versão mais nova do app."}), 400

    db_file = db.engine.url.database
    if db_file and os.path.exists(db_file):
        folder = os.path.join(os.path.dirname(db_file), "backups")
        os.makedirs(folder, exist_ok=True)
        shutil.copy2(db_file, os.path.join(folder, f"antes-de-restaurar-{_stamp()}.db"))

    try:
        for table in reversed(db.metadata.sorted_tables):
            db.session.execute(table.delete())
        restored = 0
        for table in db.metadata.sorted_tables:
            columns = {c.name: c for c in table.columns}
            for row in tables.get(table.name, []):
                values = {k: _parse(columns[k], v) for k, v in row.items() if k in columns}
                db.session.execute(table.insert().values(**values))
                restored += 1
        db.session.commit()
    except Exception as exc:  # arquivo malformado: nada é aplicado pela metade
        db.session.rollback()
        return jsonify({"error": f"Não foi possível restaurar: {exc}"}), 400
    return jsonify({"ok": True, "rows": restored})


# ---------------- CSV / Excel ----------------
TASK_HEADERS = ["ID", "Tarefa", "Categoria", "Projeto", "Prioridade", "Status", "Data", "Hora",
                "Estimado (min)", "Real (min)", "Criada em", "Concluída em"]
SESSION_HEADERS = ["Tarefa", "Início", "Fim", "Duração (min)"]


def _fmt_dt(value):
    return value.strftime("%d/%m/%Y %H:%M") if value else ""


def _task_rows():
    cats = {c.id: c.name for c in Category.query.filter_by(user_id=CURRENT_USER_ID)}
    projs = {p.id: p.name for p in Project.query.filter_by(user_id=CURRENT_USER_ID)}
    for t in Task.query.filter_by(user_id=CURRENT_USER_ID).order_by(Task.date.desc().nullslast(), Task.id.desc()):
        yield [t.id, t.name, cats.get(t.category_id, ""), projs.get(t.project_id, ""), t.priority, t.status,
               t.date.strftime("%d/%m/%Y") if t.date else "", t.time or "",
               t.estimated_minutes if t.estimated_minutes is not None else "",
               t.real_minutes if t.real_minutes is not None else "",
               _fmt_dt(t.created_at), _fmt_dt(t.completed_at)]


def _session_rows():
    names = {t.id: t.name for t in Task.query.filter_by(user_id=CURRENT_USER_ID)}
    entries = TimeEntry.query.filter(TimeEntry.task_id.in_(names.keys())).order_by(TimeEntry.start_time.desc())
    for e in entries:
        minutes = round((e.duration_seconds or 0) / 60, 1) if e.duration_seconds is not None else ""
        yield [names.get(e.task_id, ""), _fmt_dt(e.start_time), _fmt_dt(e.end_time), minutes]


def _csv(headers, rows):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")  # ';' é o que o Excel em português espera
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


@data_bp.route("/api/export/tasks.csv")
def export_tasks_csv():
    return _attachment(_csv(TASK_HEADERS, _task_rows()), "text/csv; charset=utf-8", f"meu-tempo-tarefas-{_stamp()}.csv")


@data_bp.route("/api/export/sessions.csv")
def export_sessions_csv():
    return _attachment(_csv(SESSION_HEADERS, _session_rows()), "text/csv; charset=utf-8", f"meu-tempo-sessoes-{_stamp()}.csv")


@data_bp.route("/api/export/report.xlsx")
def export_report_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tasks, sessions = list(_task_rows()), list(_session_rows())
    done = [r for r in tasks if r[5] == "concluida"]
    worked = sum(r[3] for r in sessions if isinstance(r[3], (int, float)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    summary = [
        ("Relatório Meu Tempo", ""),
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Tarefas cadastradas", len(tasks)),
        ("Tarefas concluídas", len(done)),
        ("Horas registradas", round(worked / 60, 1)),
        ("Sessões de cronômetro", len(sessions)),
    ]
    for r in summary:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22

    def sheet(title, headers, rows):
        s = wb.create_sheet(title)
        s.append(headers)
        for row in rows:
            s.append(row)
        for cell in s[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F6F6B")
            cell.alignment = Alignment(vertical="center")
        s.freeze_panes = "A2"
        for i, header in enumerate(headers, 1):
            longest = max([len(str(header))] + [len(str(r[i - 1])) for r in rows[:200]])
            s.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 46)

    sheet("Tarefas", TASK_HEADERS, tasks)
    sheet("Sessões", SESSION_HEADERS, sessions)
    out = io.BytesIO()
    wb.save(out)
    return _attachment(out.getvalue(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       f"meu-tempo-relatorio-{_stamp()}.xlsx")


# ---------------- Apresentação (primeira execução) ----------------
@data_bp.route("/api/onboarding", methods=["POST"])
def api_onboarding():
    from app.models.notebook import Notebook
    from app.models.schedule import Schedule

    data = request.get_json(silent=True) or {}
    user = db.session.get(User, CURRENT_USER_ID)
    if not data.get("skip"):
        name = (data.get("name") or "").strip()[:120]
        if name:
            user.name = name
        start, end = data.get("workday_start"), data.get("workday_end")
        valid = lambda v: isinstance(v, str) and re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", v)
        if valid(start) and valid(end) and start < end:
            user.workday_start, user.workday_end = start, end
            Schedule.query.filter_by(user_id=CURRENT_USER_ID, is_working_day=True).update(
                {"start_time": start, "end_time": end})
        notebook = (data.get("notebook") or "").strip()[:80]
        first = Notebook.query.filter_by(user_id=CURRENT_USER_ID).order_by(Notebook.id).first()
        if notebook and first:
            first.name = notebook
    user.onboarded = True
    db.session.commit()
    return jsonify({"ok": True})
